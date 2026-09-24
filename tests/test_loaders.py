import io

import pytest

from app.ingestion.loaders import EmptyDocument, UnsupportedFileType, load_document


def test_text_and_markdown():
    assert load_document("a.txt", b"hello")[0].text == "hello"
    assert "Title" in load_document("a.md", "# Title\nbody".encode())[0].text


def test_html_strips_scripts_and_tags():
    html = b"<html><head><style>x{}</style><script>alert(1)</script></head><body><h1>Refunds</h1><p>Within 30 days.</p></body></html>"
    text = load_document("p.html", html)[0].text
    assert "Refunds" in text and "Within 30 days." in text
    assert "alert" not in text and "x{}" not in text


def test_csv_rows_are_self_describing():
    text = load_document("t.csv", b"sku,price\nA1,10\nB2,20\n")[0].text
    assert "sku: A1; price: 10" in text


def test_docx_paragraphs_and_tables():
    import docx

    document = docx.Document()
    document.add_paragraph("Leave policy")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Annual"
    table.rows[0].cells[1].text = "24 days"
    buffer = io.BytesIO()
    document.save(buffer)
    text = load_document("h.docx", buffer.getvalue())[0].text
    assert "Leave policy" in text and "Annual | 24 days" in text


def test_unsupported_extension():
    with pytest.raises(UnsupportedFileType):
        load_document("x.exe", b"MZ")


def test_empty_document():
    with pytest.raises(EmptyDocument):
        load_document("blank.txt", b"   \n  ")


def test_pdf_pages_are_preserved():
    pytest.importorskip("reportlab")
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 720, "Refunds are accepted within 30 days.")
    pdf.showPage()
    pdf.drawString(72, 720, "Shipping takes five business days.")
    pdf.save()
    sections = load_document("policy.pdf", buffer.getvalue())
    assert [s.page for s in sections] == [1, 2]
    assert "Shipping" in sections[1].text
