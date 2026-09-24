"""Turn uploaded bytes into plain-text sections.

A section is the smallest unit we can cite with a location (a PDF page, or the
whole file for formats without pages). Loaders are selected by extension, so
adding a format means adding one function to LOADERS.
"""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable


class UnsupportedFileType(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


@dataclass
class Section:
    text: str
    page: int | None = None


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def load_text(data: bytes) -> list[Section]:
    return [Section(_decode(data))]


def load_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    sections = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append(Section(text, page=number))
    return sections


def load_docx(data: bytes) -> list[Section]:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return [Section("\n\n".join(parts))]


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "template", "svg"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"}

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip_depth += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip_depth:
            self.parts.append(data)


def load_html(data: bytes) -> list[Section]:
    parser = _HTMLText()
    parser.feed(_decode(data))
    lines = (" ".join(line.split()) for line in "".join(parser.parts).splitlines())
    return [Section("\n".join(line for line in lines if line))]


def load_csv(data: bytes) -> list[Section]:
    # Each row becomes "column: value" pairs so a chunk is self-describing
    # even when it no longer sits next to the header row.
    reader = csv.DictReader(io.StringIO(_decode(data)))
    rows = []
    for row in reader:
        pairs = [f"{k}: {v}" for k, v in row.items() if k and v not in (None, "")]
        if pairs:
            rows.append("; ".join(pairs))
    return [Section("\n".join(rows))]


def load_json(data: bytes) -> list[Section]:
    parsed = json.loads(_decode(data))
    return [Section(json.dumps(parsed, indent=2, ensure_ascii=False))]


LOADERS: dict[str, Callable[[bytes], list[Section]]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_text,
    ".md": load_text,
    ".markdown": load_text,
    ".rst": load_text,
    ".html": load_html,
    ".htm": load_html,
    ".csv": load_csv,
    ".json": load_json,
}

SUPPORTED_EXTENSIONS = sorted(LOADERS)


def load_document(filename: str, data: bytes) -> list[Section]:
    extension = Path(filename).suffix.lower()
    loader = LOADERS.get(extension)
    if loader is None:
        raise UnsupportedFileType(
            f"'{extension or filename}' is not supported. Use one of: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    sections = [s for s in loader(data) if s.text.strip()]
    if not sections:
        raise EmptyDocument(
            f"No extractable text found in '{filename}'. Scanned PDFs need OCR before upload."
        )
    return sections
